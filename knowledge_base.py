"""知识文件解析、向量库构建和带权限检索。"""

from __future__ import annotations

import json
import os
import shutil
import tempfile
import threading
from pathlib import Path
from typing import Any

import pymupdf
from bs4 import BeautifulSoup
from docx import Document as DocxDocument
from dotenv import load_dotenv
from langchain_chroma import Chroma
from langchain_core.documents import Document
from langchain_huggingface import HuggingFaceEmbeddings
from langchain_text_splitters import RecursiveCharacterTextSplitter
from pypdf import PdfReader
from rapidocr import RapidOCR

PROJECT_DIR = Path(__file__).resolve().parent
# 原始知识文件、知识清单和派生索引分开保存。
KNOWLEDGE_DIR = PROJECT_DIR / "knowledge"
MANIFEST_PATH = KNOWLEDGE_DIR / "day05_knowledge_manifest.json"
_configured_vector_store = (os.getenv("VECTOR_STORE_DIR") or "vector_store").strip()
VECTOR_STORE_DIR = Path(_configured_vector_store)
if not VECTOR_STORE_DIR.is_absolute():
    VECTOR_STORE_DIR = PROJECT_DIR / VECTOR_STORE_DIR
INDEX_MANIFEST_PATH = VECTOR_STORE_DIR / "index_manifest.json"
COLLECTION_NAME = "day06_company_knowledge"

_OCR_ENGINE: RapidOCR | None = None
# Embedding 模型只在进程内加载一次，避免每次检索都重新载入权重。
_EMBEDDINGS: HuggingFaceEmbeddings | None = None
_EMBEDDINGS_MODEL: str | None = None
_EMBEDDINGS_LOCK = threading.Lock()


def _hf_cache_root() -> Path:
    """返回 huggingface_hub 的本地缓存根目录。"""

    configured = (os.getenv("HF_HOME") or "").strip()
    if configured:
        return Path(configured).expanduser() / "hub"
    return Path.home() / ".cache" / "huggingface" / "hub"


def _embedding_model_is_cached(model_name: str) -> bool:
    """检查本地缓存里是否已经有该模型的快照文件。"""

    snapshot_root = _hf_cache_root() / f"models--{model_name.replace('/', '--')}" / "snapshots"
    if not snapshot_root.is_dir():
        return False
    return any(
        snapshot.is_dir() and any(snapshot.iterdir()) for snapshot in snapshot_root.iterdir()
    )


def _configure_hf_offline(model_name: str) -> bool:
    """模型已缓存时关闭 Hub 联网，避免每次加载都等待网络超时。"""

    explicit = (os.getenv("HF_HUB_OFFLINE") or "").strip().lower()
    if explicit in {"0", "false", "no"}:
        # 显式要求联网（例如首次下载模型）。
        return False
    if explicit not in {"1", "true", "yes"} and not _embedding_model_is_cached(model_name):
        return False

    os.environ["HF_HUB_OFFLINE"] = "1"
    # sentence-transformers 底层 transformers 也读取这个变量。
    os.environ.setdefault("TRANSFORMERS_OFFLINE", "1")
    return True


def _require_env(name: str) -> str:
    """读取必填配置，在项目启动时尽早暴露缺失项。"""

    value = (os.getenv(name) or "").strip()
    if not value:
        raise RuntimeError(f"缺少环境变量 {name}，请先配置 .env")
    return value


def build_embeddings() -> HuggingFaceEmbeddings:
    """创建本地 Hugging Face Embedding，构建和查询必须使用同一模型。"""

    global _EMBEDDINGS, _EMBEDDINGS_MODEL

    load_dotenv(PROJECT_DIR / ".env", override=False)
    model = _require_env("EMBEDDING_MODEL")
    with _EMBEDDINGS_LOCK:
        if _EMBEDDINGS is not None and _EMBEDDINGS_MODEL == model:
            return _EMBEDDINGS

        offline = _configure_hf_offline(model)
        try:
            embeddings = HuggingFaceEmbeddings(
                model_name=model,
                # 课堂环境统一使用 CPU，避免依赖特定显卡。
                model_kwargs={"device": "cpu"},
                # 归一化后可以使用余弦方向比较文本相似度。
                encode_kwargs={"normalize_embeddings": True},
            )
        except Exception as exc:
            hint = (
                "本地缓存中找不到该模型，且当前处于离线模式，"
                "请联网执行一次 build_index.py，或设置 HF_HUB_OFFLINE=0 后重试"
                if offline
                else "请检查网络能否访问 huggingface.co，或确认模型名是否正确"
            )
            raise RuntimeError(f"加载 Embedding 模型失败：{model}（{hint}）") from exc

        _EMBEDDINGS = embeddings
        _EMBEDDINGS_MODEL = model
        return embeddings


def _read_manifest() -> dict[str, Any]:
    if not MANIFEST_PATH.exists():
        raise FileNotFoundError(f"找不到知识清单：{MANIFEST_PATH}")
    return json.loads(MANIFEST_PATH.read_text(encoding="utf-8"))


def _read_docx(path: Path) -> str:
    """提取 DOCX 中的段落和表格文本。"""

    document = DocxDocument(str(path))
    paragraphs = [paragraph.text.strip() for paragraph in document.paragraphs]
    table_rows: list[str] = []
    for table in document.tables:
        for row in table.rows:
            cells = [cell.text.strip() for cell in row.cells]
            table_rows.append(" | ".join(cells))
    return "\n".join(item for item in [*paragraphs, *table_rows] if item)


def _read_pdf_text(path: Path) -> str:
    """提取普通文本型 PDF 中的文字。"""

    reader = PdfReader(str(path))
    return "\n".join((page.extract_text() or "").strip() for page in reader.pages).strip()


def _ocr_engine() -> RapidOCR:
    """延迟创建 OCR 引擎，处理普通文件时不需要加载它。"""

    global _OCR_ENGINE
    if _OCR_ENGINE is None:
        _OCR_ENGINE = RapidOCR()
    return _OCR_ENGINE


def _ocr_image(path: Path) -> str:
    """使用 RapidOCR 提取单张图片的可见文字。"""

    result = _ocr_engine()(str(path))
    return "\n".join(getattr(result, "txts", ()) or ()).strip()


def _ocr_pdf(path: Path) -> str:
    """把扫描 PDF 临时渲染成图片，再逐页执行 OCR。"""

    texts: list[str] = []
    with pymupdf.open(path) as pdf:
        for page_number, page in enumerate(pdf, start=1):
            pixmap = page.get_pixmap(dpi=180, alpha=False)
            with tempfile.NamedTemporaryFile(suffix=".png", delete=False) as image_file:
                image_path = Path(image_file.name)
            try:
                pixmap.save(str(image_path))
                page_text = _ocr_image(image_path)
            finally:
                image_path.unlink(missing_ok=True)
            if page_text:
                texts.append(f"[第 {page_number} 页]\n{page_text}")
    return "\n\n".join(texts)


def _extract_text(path: Path, parser_mode: str) -> str:
    """根据文件扩展名选择解析方式。"""

    suffix = path.suffix.lower()
    if suffix in {".md", ".txt"}:
        return path.read_text(encoding="utf-8").strip()
    if suffix == ".html":
        html = path.read_text(encoding="utf-8")
        return BeautifulSoup(html, "html.parser").get_text("\n", strip=True)
    if suffix == ".docx":
        return _read_docx(path)
    if suffix == ".pdf":
        text = _read_pdf_text(path)
        if parser_mode == "ocr_required" or not text:
            return _ocr_pdf(path)
        return text
    if suffix in {".png", ".jpg", ".jpeg"}:
        return _ocr_image(path)
    raise ValueError(f"暂不支持的知识文件类型：{path.name}")


def load_source_documents() -> list[Document]:
    """按清单解析原始文件，并保留来源和权限元数据。"""

    manifest = _read_manifest()
    documents: list[Document] = []

    for item in manifest["documents"]:
        # 清单中停用的文档不进入新索引。
        if item.get("status") != "active":
            continue
        source_path = KNOWLEDGE_DIR / item["path"]
        if not source_path.exists():
            raise FileNotFoundError(f"知识清单中的文件不存在：{source_path}")

        text = _extract_text(source_path, item["parser_mode"])
        if not text:
            raise ValueError(f"没有从文件中取得可索引文本：{source_path}")

        allowed_roles = set(item["allowed_roles"])
        documents.append(
            Document(
                page_content=text,
                metadata={
                    "document_id": item["document_id"],
                    "title": item["title"],
                    "source": item["path"],
                    "version": item["version"],
                    "access_scope": item["access_scope"],
                    "parser_mode": item["parser_mode"],
                    # Chroma 元数据使用标量值，分别保存每个角色是否可见。
                    "allowed_employee": "employee" in allowed_roles,
                    "allowed_hr": "hr" in allowed_roles,
                },
            )
        )
    return documents


def split_documents(documents: list[Document]) -> list[Document]:
    """切分长文档，并为每个 Chunk 生成稳定的课堂调试编号。"""

    splitter = RecursiveCharacterTextSplitter(
        chunk_size=600,
        chunk_overlap=100,
        separators=["\n\n", "\n", "。", "；", "，", " "],
    )
    chunks = splitter.split_documents(documents)
    counters: dict[str, int] = {}
    for chunk in chunks:
        document_id = str(chunk.metadata["document_id"])
        counters[document_id] = counters.get(document_id, 0) + 1
        chunk.metadata["chunk_id"] = f"{document_id}-C{counters[document_id]:03d}"
    return chunks


def _reset_vector_store() -> None:
    """只删除本项目生成的派生索引，不删除原始知识文件。"""

    expected = (PROJECT_DIR / "vector_store").resolve()
    actual = VECTOR_STORE_DIR.resolve()
    if actual != expected:
        raise RuntimeError(f"拒绝清理非预期目录：{actual}")
    if VECTOR_STORE_DIR.exists():
        shutil.rmtree(VECTOR_STORE_DIR)
    VECTOR_STORE_DIR.mkdir(parents=True, exist_ok=True)


def build_vector_store() -> dict[str, Any]:
    """重新解析全部知识文件并构建 Chroma 持久化索引。"""

    # 1. 解析原始资料并保留来源、版本和权限元数据。
    documents = load_source_documents()
    # 2. 将长文档切成适合检索的 Chunk。
    chunks = split_documents(documents)
    # 3. 创建本地 Embedding 模型。
    embeddings = build_embeddings()

    # 4. 重建的是派生索引，不会删除 knowledge/ 原始资料。
    _reset_vector_store()
    Chroma.from_documents(
        documents=chunks,
        embedding=embeddings,
        collection_name=COLLECTION_NAME,
        persist_directory=str(VECTOR_STORE_DIR),
    )

    # 5. 记录本次建库配置，供查询时检查模型是否一致。
    summary = {
        "knowledge_dataset": _read_manifest()["dataset_version"],
        "document_count": len(documents),
        "chunk_count": len(chunks),
        "embedding_model": _require_env("EMBEDDING_MODEL"),
        "collection_name": COLLECTION_NAME,
        "vector_store_path": str(VECTOR_STORE_DIR),
    }
    INDEX_MANIFEST_PATH.write_text(
        json.dumps(summary, ensure_ascii=False, indent=2),
        encoding="utf-8",
    )
    return summary


def open_vector_store() -> Chroma:
    """打开已构建的向量库，并检查模型配置是否一致。"""

    if not INDEX_MANIFEST_PATH.exists():
        raise RuntimeError("尚未构建向量库，请先运行：python build_index.py")
    summary = json.loads(INDEX_MANIFEST_PATH.read_text(encoding="utf-8"))
    current_model = _require_env("EMBEDDING_MODEL")
    # 建库和查询使用不同模型时，向量不可直接比较。
    if summary["embedding_model"] != current_model:
        raise RuntimeError("当前 Embedding 模型与索引记录不一致，请重新运行 build_index.py")
    return Chroma(
        collection_name=COLLECTION_NAME,
        embedding_function=build_embeddings(),
        persist_directory=str(VECTOR_STORE_DIR),
    )


def search_knowledge(query: str, role: str, k: int = 4) -> list[Document]:
    """在检索阶段按可信角色过滤，避免越权文档进入候选集合。"""

    if role not in {"employee", "hr"}:
        raise PermissionError("当前角色无权访问公司知识库")
    vector_store = open_vector_store()
    filters: dict[str, Any] = {f"allowed_{role}": True}
    return vector_store.similarity_search(
        query=query,
        k=k,
        # 权限过滤发生在检索阶段，无权文档不会进入候选结果。
        filter=filters,
    )
